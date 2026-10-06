from __future__ import annotations

import math
import os
import re
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Callable

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .schema import (
    CONFIG_PATH,
    MASTER_COLUMNS,
    MASTER_PATH,
    OBSERVED_MASTER_PATH,
    POLYMER_MASTER_PATH,
    REPO_ROOT,
    load_config,
)


SECTION_RE = re.compile(
    r"Mw:\s*([0-9.]+)\s*kDa,\s*f\s*=\s*([0-9.]+)%\s*,\s*([0-9.]+)\s*mg/mL",
    re.IGNORECASE,
)
JULY07_EXPERIMENT_RE = re.compile(
    r"^(?P<surfactant>.+?)\s*\|\s*(?P<salt>.+?)\s*\|\s*Run\s+(?P<run>\d+)$"
)
JULY07_SALT_RE = re.compile(
    r"^(?P<concentration>[0-9]+(?:\.[0-9]+)?)\s+mM\s+NaCl$"
)
JULY07_CONCENTRATION_RE = re.compile(
    r"^(?P<concentration>[0-9]+(?:\.[0-9]+)?)(?:-(?P<repeat>\d+))?$"
)
JULY28_SALT_RE = re.compile(
    r"^(?P<concentration>[0-9]+(?:\.[0-9]+)?)\s+mM\s+(?P<salt>[A-Za-z0-9]+)$"
)
JULY28_RUN_RE = re.compile(r"^Run\s+(?P<run>\d+)$")
NO_SURFACTANT_NACL_SEQUENCE_MM = (0.1, 1.0, 5.0, 10.0, 20.0, 50.0, 100.0)
CONFIRMED_COPIED_SOURCE_ROWS = {
    ("Hofmiester series", 17),
    ("Hofmiester series", 36),
}
LEGACY_NONE_MARKERS = {"N/A", "NA", "NONE", "-"}
TARGET_POLICIES = {"observed", "polymer"}
TRANSITION_ATTRIBUTION_COLUMN = "_transition_attribution"
INGEST_COLUMNS = [*MASTER_COLUMNS, TRANSITION_ATTRIBUTION_COLUMN]
EXPECTED_TRANSITION_COUNTS = {
    "observed": {"LCST": 927, "UCST": 56, "NONE": 272},
    "polymer": {"LCST": 913, "UCST": 43, "NONE": 299},
}


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip()
    if not text or text.upper() in LEGACY_NONE_MARKERS:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _legacy_trial_outcomes(
    values: tuple[object, ...] | list[object],
) -> list[tuple[str, float | None]]:
    outcomes: list[tuple[str, float | None]] = []
    for value in values:
        temperature = _number(value)
        if temperature is not None:
            outcomes.append(("LCST", temperature))
        elif isinstance(value, str) and value.strip().upper() in LEGACY_NONE_MARKERS:
            outcomes.append(("NONE", None))
    return outcomes


def _molar(value: object) -> float | None:
    direct = _number(value)
    if direct is not None:
        return direct
    if not isinstance(value, str):
        return None
    match = re.search(r"([0-9]*\.?[0-9]+)\s*([munp]?m)\b", value, re.IGNORECASE)
    if not match:
        return None
    amount = float(match.group(1))
    unit = match.group(2)
    if unit == "M":
        return amount
    factors = {"m": 1.0, "mm": 1e-3, "um": 1e-6, "nm": 1e-9, "pm": 1e-12}
    return amount * factors[unit.lower()]


def _cell(rows: list[tuple[object, ...]], row: int, column: int) -> object:
    if row >= len(rows) or column >= len(rows[row]):
        return None
    return rows[row][column]


def _master_record(
    *,
    workbook: str,
    sheet: str,
    polymer_name: str,
    polymer_mw_kda: float | None,
    polymer_functionalization_percent: float | None,
    polymer_concentration_mg_ml: float | None,
    buffer: str,
    heating_rate_c_per_min: float | None,
    additive_name: str,
    additive_concentration_mm: float,
    salt_name: str,
    salt_concentration_mm: float,
    transition_type: str,
    transition_temperature_c: float | None,
    transition_attribution: str = "polymer",
) -> dict[str, object]:
    if polymer_name != "Dex-MA":
        polymer_name = "No polymer"
        polymer_mw_kda = None
        polymer_functionalization_percent = None
        polymer_concentration_mg_ml = 0.0
        if transition_type != "NONE":
            transition_attribution = "no_polymer"
    if additive_concentration_mm <= 0:
        additive_name = "No additive"
        additive_concentration_mm = 0.0
    if salt_concentration_mm <= 0:
        salt_name = "No salt"
        salt_concentration_mm = 0.0
    if transition_type == "NONE":
        transition_temperature_c = None
        transition_attribution = "none"
    return {
        "source_workbook": workbook,
        "source_sheet": sheet,
        "polymer_name": polymer_name,
        "polymer_mw_kda": polymer_mw_kda,
        "polymer_functionalization_percent": polymer_functionalization_percent,
        "polymer_concentration_mg_ml": polymer_concentration_mg_ml,
        "buffer": buffer,
        "heating_rate_c_per_min": heating_rate_c_per_min,
        "additive_name": additive_name,
        "additive_concentration_mM": additive_concentration_mm,
        "salt_name": salt_name,
        "salt_concentration_mM": salt_concentration_mm,
        "transition_type": transition_type,
        "transition_temperature_c": transition_temperature_c,
        TRANSITION_ATTRIBUTION_COLUMN: transition_attribution,
    }


def _add_legacy_trials(
    records: list[dict[str, object]],
    *,
    workbook: str,
    sheet: str,
    polymer_mw_kda: float | None,
    polymer_functionalization_percent: float | None,
    polymer_concentration_mg_ml: float | None,
    additive_name: str = "No additive",
    additive_concentration_m: float = 0.0,
    salt_name: str = "No salt",
    salt_concentration_m: float = 0.0,
    heating_rate: float | None = None,
    trials: tuple[object, ...] | list[object],
) -> None:
    for transition_type, temperature in _legacy_trial_outcomes(trials):
        records.append(
            _master_record(
                workbook=workbook,
                sheet=sheet,
                polymer_name="Dex-MA",
                polymer_mw_kda=polymer_mw_kda,
                polymer_functionalization_percent=polymer_functionalization_percent,
                polymer_concentration_mg_ml=polymer_concentration_mg_ml,
                buffer="PBS",
                heating_rate_c_per_min=heating_rate,
                additive_name=additive_name,
                additive_concentration_mm=additive_concentration_m * 1000.0,
                salt_name=salt_name,
                salt_concentration_mm=salt_concentration_m * 1000.0,
                transition_type=transition_type,
                transition_temperature_c=temperature,
            )
        )


def _parse_legacy_standard_sheets(workbook_path: Path, wb) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    workbook = workbook_path.name

    rows = list(wb["LCST for 86kDa Dex-MA"].iter_rows(values_only=True))
    for excel_row in range(4, 7):
        idx = excel_row - 1
        _add_legacy_trials(
            records,
            workbook=workbook,
            sheet="LCST for 86kDa Dex-MA",
            polymer_mw_kda=86.0,
            polymer_functionalization_percent=_number(_cell(rows, idx, 1)),
            polymer_concentration_mg_ml=10.0,
            trials=list(rows[idx][2:]),
        )
    for excel_row in range(11, 19):
        idx = excel_row - 1
        _add_legacy_trials(
            records,
            workbook=workbook,
            sheet="LCST for 86kDa Dex-MA",
            polymer_mw_kda=86.0,
            polymer_functionalization_percent=88.0,
            polymer_concentration_mg_ml=_number(_cell(rows, idx, 1)),
            trials=list(rows[idx][2:]),
        )

    rows = list(wb["Hofmiester series"].iter_rows(values_only=True))
    ranges = [
        (range(5, 13), 0.1, None),
        (range(17, 22), 0.1, None),
        (range(26, 30), None, "NaCl"),
        (range(35, 37), None, "Na2SO4"),
    ]
    for excel_rows, fixed_concentration, fixed_salt in ranges:
        for excel_row in excel_rows:
            if ("Hofmiester series", excel_row) in CONFIRMED_COPIED_SOURCE_ROWS:
                continue
            idx = excel_row - 1
            salt = fixed_salt or str(_cell(rows, idx, 1)).strip()
            concentration = fixed_concentration if fixed_concentration is not None else _molar(_cell(rows, idx, 1))
            _add_legacy_trials(
                records,
                workbook=workbook,
                sheet="Hofmiester series",
                polymer_mw_kda=86.0,
                polymer_functionalization_percent=88.0,
                # Erfan Moaseri confirmed with Linqing Li on 2026-09-22 that
                # every formulation in this sheet used 10 mg/mL Dex-MA (88%).
                # See data/metadata/hofmeister_concentration_20260922.json.
                polymer_concentration_mg_ml=10.0,
                salt_name=salt,
                salt_concentration_m=concentration or 0.0,
                trials=list(rows[idx][2:]),
            )

    rows = list(wb["LCST for Diff Mw of Dextran"].iter_rows(values_only=True))
    for molecular_weight, excel_rows in [
        (86.0, range(5, 11)),
        (250.0, range(17, 23)),
        (500.0, range(29, 35)),
    ]:
        for excel_row in excel_rows:
            idx = excel_row - 1
            _add_legacy_trials(
                records,
                workbook=workbook,
                sheet="LCST for Diff Mw of Dextran",
                polymer_mw_kda=molecular_weight,
                polymer_functionalization_percent=70.0,
                polymer_concentration_mg_ml=_number(_cell(rows, idx, 1)),
                trials=list(rows[idx][2:]),
            )

    rows = list(wb["LCST of DexMA in Urea"].iter_rows(values_only=True))
    for excel_row in range(5, 8):
        idx = excel_row - 1
        _add_legacy_trials(
            records,
            workbook=workbook,
            sheet="LCST of DexMA in Urea",
            polymer_mw_kda=86.0,
            polymer_functionalization_percent=88.0,
            # Figure 4b explicitly reports 1 mg/mL Dex-MA (86 kDa, 88%)
            # for 1, 2, and 4 M urea. This is not a default for other sheets.
            # https://www.nature.com/articles/s41598-023-41947-z/figures/4
            polymer_concentration_mg_ml=1.0,
            additive_name="Urea",
            additive_concentration_m=_molar(_cell(rows, idx, 1)) or 0.0,
            trials=list(rows[idx][2:]),
        )

    for sheet, additive in [
        ("LCST with CTAB", "CTAB"),
        ("LCST with SDS", "SDS"),
        ("LCST with CHAPS", "CHAPS"),
        ("LCST with Pluronic", "Pluronic F127"),
    ]:
        rows = list(wb[sheet].iter_rows(values_only=True))
        row_idx = 0
        while row_idx < len(rows):
            title = next(
                (
                    value
                    for value in rows[row_idx]
                    if isinstance(value, str) and "Transition temperature values" in value
                ),
                None,
            )
            match = SECTION_RE.search(title) if title else None
            if not match:
                row_idx += 1
                continue
            molecular_weight, functionalization, polymer_concentration = map(float, match.groups())
            header_idx = None
            concentration_column = None
            for scan_idx in range(row_idx + 1, min(row_idx + 6, len(rows))):
                for column, value in enumerate(rows[scan_idx]):
                    if isinstance(value, str) and value.strip().lower().startswith("concentration of"):
                        header_idx, concentration_column = scan_idx, column
                        break
                if header_idx is not None:
                    break
            if header_idx is None or concentration_column is None:
                row_idx += 1
                continue
            data_idx = header_idx + 1
            while data_idx < len(rows):
                if any(
                    isinstance(value, str) and "Transition temperature values" in value
                    for value in rows[data_idx]
                ):
                    break
                concentration = _molar(_cell(rows, data_idx, concentration_column))
                if concentration is not None:
                    _add_legacy_trials(
                        records,
                        workbook=workbook,
                        sheet=sheet,
                        polymer_mw_kda=molecular_weight,
                        polymer_functionalization_percent=functionalization,
                        polymer_concentration_mg_ml=polymer_concentration,
                        additive_name=additive,
                        additive_concentration_m=concentration,
                        trials=list(rows[data_idx][concentration_column + 1 :]),
                    )
                data_idx += 1
            row_idx = data_idx
    return records


def _parse_heating_sheet(workbook_path: Path, wb) -> list[dict[str, object]]:
    sheet = "LCST - Different HeatingRate"
    rows = list(wb[sheet].iter_rows(values_only=True))
    records: list[dict[str, object]] = []
    heating_rate: float | None = None
    for excel_row in range(7, 35):
        idx = excel_row - 1
        heating_rate = _number(_cell(rows, idx, 2)) or heating_rate
        concentration = _number(_cell(rows, idx, 3))
        _add_legacy_trials(
            records,
            workbook=workbook_path.name,
            sheet=sheet,
            polymer_mw_kda=86.0,
            polymer_functionalization_percent=88.0,
            polymer_concentration_mg_ml=concentration,
            heating_rate=heating_rate,
            trials=list(rows[idx][4:7]),
        )
    for excel_row in range(48, 55):
        idx = excel_row - 1
        _add_legacy_trials(
            records,
            workbook=workbook_path.name,
            sheet=sheet,
            polymer_mw_kda=86.0,
            polymer_functionalization_percent=88.0,
            polymer_concentration_mg_ml=_number(_cell(rows, idx, 2)),
            heating_rate=1.0,
            trials=list(rows[idx][3:7]),
        )
    return records


def parse_legacy(workbook_path: Path, source_config: dict[str, Any]) -> pd.DataFrame:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    configured_sheets = set(source_config["sheets"])
    missing_sheets = configured_sheets.difference(wb.sheetnames)
    if missing_sheets:
        raise ValueError(f"{workbook_path.name}: missing configured sheets {sorted(missing_sheets)}")
    records = _parse_legacy_standard_sheets(workbook_path, wb)
    records.extend(_parse_heating_sheet(workbook_path, wb))
    return pd.DataFrame.from_records(records, columns=INGEST_COLUMNS)


def _july07_concentration(value: Any) -> float:
    if isinstance(value, (int, float, np.integer, np.floating)) and not pd.isna(value):
        return float(value)
    match = JULY07_CONCENTRATION_RE.fullmatch(str(value).strip())
    if not match:
        raise ValueError(f"Unrecognized July 7 surfactant concentration: {value!r}")
    return float(match.group("concentration"))


def parse_july07(workbook_path: Path, source_config: dict[str, Any]) -> pd.DataFrame:
    sheet = str(source_config["sheet"])
    columns = source_config["columns"]
    source = pd.read_excel(workbook_path, sheet_name=sheet)
    no_surfactant_positions: defaultdict[str, int] = defaultdict(int)
    records: list[dict[str, object]] = []
    for _, row in source.iterrows():
        raw_condition = str(row[columns["condition"]]).strip()
        match = JULY07_EXPERIMENT_RE.fullmatch(raw_condition)
        if not match:
            raise ValueError(f"Unrecognized July 7 condition: {raw_condition!r}")
        raw_additive = match.group("surfactant").strip()
        raw_salt = match.group("salt").strip()
        if raw_additive == "No surfactant":
            position = no_surfactant_positions[raw_condition]
            if position >= len(NO_SURFACTANT_NACL_SEQUENCE_MM):
                raise ValueError(f"Too many no-surfactant rows for {raw_condition!r}")
            additive, additive_mm = "No additive", 0.0
            salt_mm = NO_SURFACTANT_NACL_SEQUENCE_MM[position]
            no_surfactant_positions[raw_condition] += 1
        else:
            additive = "Pluronic F127" if raw_additive == "Pluronic" else raw_additive
            additive_mm = _july07_concentration(row[columns["additive_concentration_mM"]])
            if raw_salt in {"0 mM NaCl", "No added NaCl"}:
                salt_mm = 0.0
            else:
                salt_match = JULY07_SALT_RE.fullmatch(raw_salt)
                if not salt_match:
                    raise ValueError(f"Unrecognized July 7 salt: {raw_salt!r}")
                salt_mm = float(salt_match.group("concentration"))
        raw_transition = str(row[columns["transition_type"]]).strip()
        transition = "NONE" if raw_transition == "-" else raw_transition
        if transition not in {"LCST", "UCST", "NONE"}:
            raise ValueError(f"Unrecognized July 7 transition: {raw_transition!r}")
        records.append(
            _master_record(
                workbook=workbook_path.name,
                sheet=sheet,
                polymer_name=str(source_config["polymer_name"]),
                polymer_mw_kda=float(source_config["polymer_mw_kda"]),
                polymer_functionalization_percent=float(source_config["polymer_functionalization_percent"]),
                polymer_concentration_mg_ml=float(source_config["polymer_concentration_mg_ml"]),
                buffer=str(source_config["buffer"]),
                heating_rate_c_per_min=float(source_config["heating_rate_c_per_min"]),
                additive_name=additive,
                additive_concentration_mm=additive_mm,
                salt_name="NaCl",
                salt_concentration_mm=salt_mm,
                transition_type=transition,
                transition_temperature_c=_number(row[columns["transition_temperature_c"]]),
            )
        )
    return pd.DataFrame.from_records(records, columns=INGEST_COLUMNS)


def _parse_july28_condition(raw_condition: str) -> dict[str, object]:
    tokens = [token.strip() for token in raw_condition.split("|")]
    if len(tokens) < 2 or tokens[0] != "SDS":
        raise ValueError(f"Unrecognized July 28 condition: {raw_condition!r}")
    salt_match = JULY28_SALT_RE.fullmatch(tokens[1])
    if not salt_match:
        raise ValueError(f"Unrecognized July 28 salt: {tokens[1]!r}")
    modifiers = tokens[2:]
    for modifier in modifiers:
        if modifier in {"No DexMA", "DI Water"} or JULY28_RUN_RE.fullmatch(modifier):
            continue
        raise ValueError(f"Unrecognized July 28 modifier: {modifier!r}")
    return {
        "salt_name": salt_match.group("salt"),
        "salt_concentration_mM": float(salt_match.group("concentration")),
        "has_polymer": "No DexMA" not in modifiers,
        "buffer": "DI Water" if "DI Water" in modifiers else "PBS",
    }


def parse_july28(workbook_path: Path, source_config: dict[str, Any]) -> pd.DataFrame:
    sheet = str(source_config["sheet"])
    columns = source_config["columns"]
    source = pd.read_excel(workbook_path, sheet_name=sheet, keep_default_na=False)
    records: list[dict[str, object]] = []
    for _, row in source.iterrows():
        parsed = _parse_july28_condition(str(row[columns["condition"]]).strip())
        has_polymer = bool(parsed["has_polymer"])
        additive_mm = float(row[columns["additive_concentration_mM"]])
        raw_transition = str(row[columns["transition_type"]]).strip()
        transition = "NONE" if raw_transition.upper() in LEGACY_NONE_MARKERS else raw_transition
        if transition not in {"LCST", "UCST", "NONE"}:
            raise ValueError(f"Unrecognized July 28 transition: {raw_transition!r}")
        records.append(
            _master_record(
                workbook=workbook_path.name,
                sheet=sheet,
                polymer_name=str(source_config["polymer_name"]) if has_polymer else "No polymer",
                polymer_mw_kda=float(source_config["polymer_mw_kda"]) if has_polymer else None,
                polymer_functionalization_percent=float(source_config["polymer_functionalization_percent"]) if has_polymer else None,
                polymer_concentration_mg_ml=float(source_config["polymer_concentration_mg_ml"]) if has_polymer else 0.0,
                buffer=str(parsed["buffer"]),
                heating_rate_c_per_min=float(source_config["heating_rate_c_per_min"]),
                additive_name="SDS",
                additive_concentration_mm=additive_mm,
                salt_name=str(parsed["salt_name"]),
                salt_concentration_mm=float(parsed["salt_concentration_mM"]),
                transition_type=transition,
                transition_temperature_c=_number(row[columns["transition_temperature_c"]]),
            )
        )
    return pd.DataFrame.from_records(records, columns=INGEST_COLUMNS)


def parse_august(workbook_path: Path, source_config: dict[str, Any]) -> pd.DataFrame:
    sheet = str(source_config["sheet"])
    columns = source_config["columns"]
    source = pd.read_excel(workbook_path, sheet_name=sheet, keep_default_na=False)
    records: list[dict[str, object]] = []
    for _, row in source.iterrows():
        has_polymer = str(row[columns["polymer_name"]]).strip() == "Dex-MA"
        reported = " ".join(str(row[columns["transition_type"]]).split())
        if reported.upper() in LEGACY_NONE_MARKERS:
            transition = "NONE"
            attribution = "none"
        elif reported in {"LCST", "UCST"}:
            transition = reported
            attribution = "polymer"
        elif reported in {"LCST (solvent)", "UCST (solvent)"}:
            transition = reported.split()[0]
            attribution = "solvent"
        else:
            raise ValueError(f"Unrecognized August transition: {reported!r}")
        records.append(
            _master_record(
                workbook=workbook_path.name,
                sheet=sheet,
                polymer_name="Dex-MA" if has_polymer else "No polymer",
                polymer_mw_kda=_number(row[columns["polymer_mw_kda"]]) if has_polymer else None,
                polymer_functionalization_percent=_number(row[columns["polymer_functionalization_percent"]]) if has_polymer else None,
                polymer_concentration_mg_ml=_number(row[columns["polymer_concentration_mg_ml"]]) if has_polymer else 0.0,
                buffer=str(row[columns["buffer"]]).strip(),
                heating_rate_c_per_min=_number(row[columns["heating_rate_c_per_min"]]),
                additive_name=str(row[columns["additive_name"]]).strip(),
                additive_concentration_mm=float(row[columns["additive_concentration_mM"]]),
                salt_name=str(row[columns["salt_name"]]).strip(),
                salt_concentration_mm=float(row[columns["salt_concentration_mM"]]),
                transition_type=transition,
                transition_temperature_c=_number(row[columns["transition_temperature_c"]]),
                transition_attribution=attribution,
            )
        )
    return pd.DataFrame.from_records(records, columns=INGEST_COLUMNS)


ADAPTERS: dict[str, Callable[[Path, dict[str, Any]], pd.DataFrame]] = {
    "legacy": parse_legacy,
    "july07": parse_july07,
    "july28": parse_july28,
    "august": parse_august,
}


def validate_master(
    master: pd.DataFrame,
    expected_source_counts: dict[str, int],
    target_policy: str,
) -> None:
    if target_policy not in TARGET_POLICIES:
        raise ValueError(f"Unknown target policy: {target_policy!r}")
    if list(master.columns) != MASTER_COLUMNS:
        raise ValueError("Master columns do not match the canonical schema")
    if len(master) != 1255:
        raise ValueError(f"Expected 1,255 master rows; found {len(master)}")
    source_counts = master["source_workbook"].value_counts().to_dict()
    if source_counts != expected_source_counts:
        raise ValueError(f"Unexpected source counts: {source_counts}")
    transition_counts = master["transition_type"].value_counts().to_dict()
    expected_transitions = EXPECTED_TRANSITION_COUNTS[target_policy]
    if transition_counts != expected_transitions:
        raise ValueError(f"Unexpected transition counts: {transition_counts}")
    none_rows = master["transition_type"].eq("NONE")
    if master.loc[none_rows, "transition_temperature_c"].notna().any():
        raise ValueError("NONE rows must not have transition temperatures")
    if master.loc[~none_rows, "transition_temperature_c"].isna().any():
        raise ValueError("LCST and UCST rows must have transition temperatures")
    no_polymer = master["polymer_name"].eq("No polymer")
    no_polymer_counts = master.loc[no_polymer, "transition_type"].value_counts().to_dict()
    if target_policy == "polymer" and no_polymer_counts != {"NONE": 48}:
        raise ValueError(f"Unexpected polymer-policy control counts: {no_polymer_counts}")
    if target_policy == "observed" and no_polymer_counts != {"NONE": 33, "UCST": 11, "LCST": 4}:
        raise ValueError(f"Unexpected observed control counts: {no_polymer_counts}")


def _parsed_measurements(config_path: Path) -> tuple[pd.DataFrame, dict[str, int]]:
    config = load_config(config_path)
    defaults = config["parsing_defaults"]
    frames: list[pd.DataFrame] = []
    expected_source_counts: dict[str, int] = {}
    for configured_source in config["sources"]:
        source = {**defaults, **configured_source}
        path = REPO_ROOT / source["path"]
        if not path.exists():
            raise FileNotFoundError(path)
        adapter_name = str(source["adapter"])
        if adapter_name not in ADAPTERS:
            raise ValueError(f"Unknown source adapter: {adapter_name}")
        frame = ADAPTERS[adapter_name](path, source)
        expected_rows = int(source["expected_rows"])
        if len(frame) != expected_rows:
            raise ValueError(f"{path.name}: expected {expected_rows} rows; found {len(frame)}")
        frames.append(frame)
        expected_source_counts[path.name] = expected_rows
    parsed = pd.concat(frames, ignore_index=True)[INGEST_COLUMNS]
    return parsed, expected_source_counts


def _target_view(parsed: pd.DataFrame, target_policy: str) -> pd.DataFrame:
    if target_policy not in TARGET_POLICIES:
        raise ValueError(f"Unknown target policy: {target_policy!r}")
    view = parsed.copy()
    if target_policy == "polymer":
        assumed_none = (
            view["polymer_name"].ne("Dex-MA")
            | view[TRANSITION_ATTRIBUTION_COLUMN].eq("solvent")
        )
        view.loc[assumed_none, "transition_type"] = "NONE"
        view.loc[assumed_none, "transition_temperature_c"] = np.nan
    return view[MASTER_COLUMNS]


def _validate_policy_pair(observed: pd.DataFrame, polymer: pd.DataFrame) -> None:
    target_columns = ["transition_type", "transition_temperature_c"]
    shared_columns = [column for column in MASTER_COLUMNS if column not in target_columns]
    pd.testing.assert_frame_equal(
        observed[shared_columns],
        polymer[shared_columns],
        check_dtype=True,
        check_exact=True,
    )
    type_differs = observed["transition_type"].ne(polymer["transition_type"])
    observed_temperatures = observed["transition_temperature_c"]
    polymer_temperatures = polymer["transition_temperature_c"]
    temperature_differs = ~(
        observed_temperatures.eq(polymer_temperatures)
        | (observed_temperatures.isna() & polymer_temperatures.isna())
    )
    target_differs = type_differs | temperature_differs
    if int(target_differs.sum()) != 27:
        raise ValueError(f"Expected 27 policy-dependent measurements; found {int(target_differs.sum())}")
    if not polymer.loc[target_differs, "transition_type"].eq("NONE").all():
        raise ValueError("Every polymer-policy difference must map to NONE")
    if polymer.loc[target_differs, "transition_temperature_c"].notna().any():
        raise ValueError("Polymer-policy NONE targets must not retain temperatures")


def build_masters(config_path: Path = CONFIG_PATH) -> dict[str, pd.DataFrame]:
    parsed, expected_source_counts = _parsed_measurements(config_path)
    masters = {
        policy: _target_view(parsed, policy)
        for policy in ("observed", "polymer")
    }
    for policy, master in masters.items():
        validate_master(master, expected_source_counts, policy)
    _validate_policy_pair(masters["observed"], masters["polymer"])
    return masters


def build_master(
    config_path: Path = CONFIG_PATH,
    target_policy: str = "observed",
) -> pd.DataFrame:
    return build_masters(config_path)[target_policy]


def _write_master(master: pd.DataFrame, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        newline="",
        suffix=".csv",
        prefix=f"{output_path.stem}_",
        dir=output_path.parent,
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
        master.to_csv(handle, index=False, lineterminator="\n")
    try:
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
    return output_path


def rebuild_master(
    output_path: Path = MASTER_PATH,
    config_path: Path = CONFIG_PATH,
    target_policy: str = "observed",
) -> Path:
    master = build_master(config_path, target_policy)
    return _write_master(master, output_path)


def rebuild_masters(
    observed_output_path: Path = OBSERVED_MASTER_PATH,
    polymer_output_path: Path = POLYMER_MASTER_PATH,
    config_path: Path = CONFIG_PATH,
) -> dict[str, Path]:
    if observed_output_path.resolve() == polymer_output_path.resolve():
        raise ValueError("Observed and polymer masters must use different output paths")
    masters = build_masters(config_path)
    return {
        "observed": _write_master(masters["observed"], observed_output_path),
        "polymer": _write_master(masters["polymer"], polymer_output_path),
    }
